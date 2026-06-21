"""Excel 结果回写副本生成。"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def write_results_copy(
    excel_path: str | Path,
    sheet_name: str,
    case_summaries: list[dict[str, str]],
    output_path: str | Path,
) -> Path:
    """把运行结果写入 Excel 副本，不覆盖原始工作簿。"""

    try:
        import openpyxl  # type: ignore
    except ModuleNotFoundError as exc:
        raise RuntimeError("openpyxl is required to write Excel result copies.") from exc

    source = Path(excel_path)
    output = Path(output_path)
    workbook = openpyxl.load_workbook(source)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = _headers(sheet)
    result_col = _required_column(headers, "测试结果")
    note_col = _required_column(headers, "备注")
    case_col = _required_column(headers, "TC_用例名称")
    summaries = {_base_case_id(item["case_id"]): item for item in case_summaries}
    summaries.update({item["case_id"]: item for item in case_summaries})

    for row_index in range(2, sheet.max_row + 1):
        case_id = str(sheet.cell(row=row_index, column=case_col).value or "").strip()
        summary = summaries.get(case_id) or summaries.get(f"{case_id}__row_{row_index}")
        if not summary:
            continue
        sheet.cell(row=row_index, column=result_col).value = summary.get("run_status", "")
        sheet.cell(row=row_index, column=note_col).value = _note(summary)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _headers(sheet: Any) -> dict[str, int]:
    return {str(sheet.cell(row=1, column=col).value or "").strip(): col for col in range(1, sheet.max_column + 1)}


def _required_column(headers: dict[str, int], name: str) -> int:
    if name not in headers:
        raise ValueError(f"Missing required Excel header: {name}")
    return headers[name]


def _base_case_id(case_id: str) -> str:
    return case_id.split("__row_", 1)[0]


def _note(summary: dict[str, str]) -> str:
    parts = [
        summary.get("summary", ""),
        f"证据: {summary.get('evidence_dir', '')}" if summary.get("evidence_dir") else "",
        f"人工复核原因: {summary.get('manual_review_reason', '')}" if summary.get("manual_review_reason") else "",
    ]
    return "；".join(part for part in parts if part)
