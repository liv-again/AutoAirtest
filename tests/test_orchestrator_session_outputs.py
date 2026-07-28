import json

from autoairtest import orchestrator
from autoairtest.agents.planner import RuleBasedPlanner as RealRuleBasedPlanner
from autoairtest.models import ActionResult, ActionStatus, NaturalLanguageTestCase, PreliminaryJudgment, PreliminaryStatus


def test_run_offline_writes_session_metadata_steps_crashes_and_state_graph(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_session",
        internal_id="TC_session",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="跳转校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False, "session_name": "smoke"},
            "input": {"case_filter": ""},
        }
    )

    assert run_dir.name.endswith("_smoke")
    session_meta = json.loads((run_dir / "session_meta.json").read_text(encoding="utf-8"))
    assert session_meta["workflow"] == "excel_case_run"
    assert session_meta["case_count"] == 1
    assert (run_dir / "config.resolved.json").exists()
    assert (run_dir / "crashes.jsonl").read_text(encoding="utf-8") == ""
    assert json.loads((run_dir / "state_graph.json").read_text(encoding="utf-8")) == {"pages": {}, "edges": []}
    steps = [json.loads(line) for line in (run_dir / "steps.jsonl").read_text(encoding="utf-8").splitlines()]
    assert steps
    assert steps[0]["case_id"] == "TC_session"
    assert steps[0]["crash_count"] == 0


def test_run_offline_records_log_collector_crashes(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_crash",
        internal_id="TC_crash",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="跳转校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    class FakeLogCollector:
        def __init__(self, *args, **kwargs):
            self.calls = 0

        def get_recent_crashes(self, package="", lines=300):
            self.calls += 1
            if self.calls == 1:
                return {
                    "status": "success",
                    "crash_count": 1,
                    "crashes": [
                        {
                            "signature_id": "abc123",
                            "kind": "java",
                            "exception_class": "java.lang.NullPointerException",
                            "top_frames_normalized": ["com.example.LoginActivity.onClick"],
                            "process": package,
                            "source": "logcat",
                        }
                    ],
                }
            return {"status": "success", "crash_count": 0, "crashes": []}

    monkeypatch.setattr(orchestrator, "LogCollector", FakeLogCollector)

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "app": {"package": "com.example.securities"},
            "input": {"case_filter": ""},
        }
    )

    steps = [json.loads(line) for line in (run_dir / "steps.jsonl").read_text(encoding="utf-8").splitlines()]
    crashes = [json.loads(line) for line in (run_dir / "crashes.jsonl").read_text(encoding="utf-8").splitlines()]
    crash_refs = json.loads((run_dir / "cases" / "TC_crash" / "crash_refs.json").read_text(encoding="utf-8"))
    assert steps[0]["crash_count"] == 1
    assert crashes[0]["crash_id"] == "c1"
    assert crashes[0]["case_id"] == "TC_crash"
    assert crashes[0]["step_index"] == steps[0]["index"]
    assert crashes[0]["signature"]["signature_id"] == "abc123"
    assert crash_refs == [
        {
            "crash_id": "c1",
            "step_index": steps[0]["index"],
            "action_id": steps[0]["action_id"],
            "signature_id": "abc123",
        }
    ]


def test_run_offline_deduplicates_crash_signatures(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_crash_dedup",
        internal_id="TC_crash_dedup",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="崩溃去重",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：进入国内指数并点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    class FakeLogCollector:
        def __init__(self, *args, **kwargs):
            pass

        def get_recent_crashes(self, package="", lines=300):
            return {
                "status": "success",
                "crash_count": 1,
                "crashes": [
                    {
                        "signature_id": "same-crash",
                        "kind": "java",
                        "exception_class": "java.lang.NullPointerException",
                        "top_frames_normalized": ["com.example.LoginActivity.onClick"],
                        "process": package,
                        "source": "logcat",
                    }
                ],
            }

    monkeypatch.setattr(orchestrator, "LogCollector", FakeLogCollector)

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "app": {"package": "com.example.securities"},
            "input": {"case_filter": ""},
        }
    )

    crashes = [json.loads(line) for line in (run_dir / "crashes.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(crashes) == 1
    assert crashes[0]["signature"]["signature_id"] == "same-crash"
    assert crashes[0]["step_index"] == 1


def test_run_offline_passes_manual_review_reason_to_html_report(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_manual",
        internal_id="TC_manual",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="数据校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="数据显示正确",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": True},
            "logs": {"enable_capture": False},
            "input": {"case_filter": ""},
        }
    )

    html = (run_dir / "report.html").read_text(encoding="utf-8")
    assert "data_correctness" in html


def test_run_offline_writes_excel_result_copy_when_enabled(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook

    excel_path = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求测试报告"
    sheet.append(["TC_用例名称", "测试结果", "备注"])
    sheet.append(["TC_excel", "", ""])
    workbook.save(excel_path)

    case = NaturalLanguageTestCase(
        case_id="TC_excel",
        internal_id="TC_excel",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="数据校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="数据显示正确",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path / "runs"), "generate_html": False, "write_back_excel": True},
            "input": {"excel_path": str(excel_path), "sheet_name": "需求测试报告", "case_filter": ""},
            "logs": {"enable_capture": False},
        }
    )

    output = run_dir / "test-cases-with-results.xlsx"
    assert output.exists()
    copied = load_workbook(output)["需求测试报告"]
    assert copied.cell(row=2, column=2).value == "blocked"
    assert "data_correctness" in copied.cell(row=2, column=3).value


def test_run_offline_uses_device_workflow_when_execution_mode_is_device(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_device",
        internal_id="TC_device",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="跳转校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    class FakeDeviceWorkflow:
        def execute_plan(self, plan, case_dir):
            marker = case_dir / "device-workflow-called.txt"
            marker.write_text(plan.case_id, encoding="utf-8")
            return []

    workflow_kwargs = {}

    def fake_device_workflow(*args, **kwargs):
        workflow_kwargs.update(kwargs)
        return FakeDeviceWorkflow()

    monkeypatch.setattr(orchestrator, "DeviceWorkflow", fake_device_workflow)

    run_dir = orchestrator.run_offline(
        {
            "execution": {"mode": "device", "retry": {"poco_dump_max_attempts": 5}},
            "app": {"package": "com.example.securities", "activity": ".MainActivity"},
            "evidence": {"redact_sensitive_text": False, "sensitive_keywords": ["手机号"]},
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "input": {"case_filter": ""},
        }
    )

    assert (run_dir / "cases" / "TC_device" / "device-workflow-called.txt").read_text(encoding="utf-8") == "TC_device"
    assert workflow_kwargs["retry_config"]["poco_dump_max_attempts"] == 5
    assert workflow_kwargs["app_config"]["package"] == "com.example.securities"
    assert workflow_kwargs["app_config"]["activity"] == ".MainActivity"
    assert workflow_kwargs["evidence_config"]["redact_sensitive_text"] is False
    assert workflow_kwargs["evidence_config"]["sensitive_keywords"] == ["手机号"]


def test_run_offline_recollects_evidence_for_verification_gap(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_recollect",
        internal_id="TC_recollect",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="顺序校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="依次展示上证、深证、创业板",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])
    monkeypatch.setattr(
        orchestrator,
        "verify_goals",
        lambda goals, evidence: [
            PreliminaryJudgment(
                goal_id="g_order",
                preliminary_status=PreliminaryStatus.MANUAL_REQUIRED,
                confidence=0.0,
                basis="Order evidence is incomplete.",
                evidence_files=[],
                human_review_required=True,
                review_reason="verification_evidence_gap",
                manual_review_reason="verification_evidence_gap",
            )
        ],
    )

    class FakeEvidenceRecollector:
        def __init__(self, max_attempts=2):
            self.max_attempts = max_attempts

        def recollect(self, goal_id, case_dir, attempt):
            return {
                "status": "collected",
                "goal_id": goal_id,
                "attempt": attempt,
                "action": "refresh_current_screen_evidence" if attempt == 1 else "visibility_adjustment",
                "action_risk_level": "low",
            }

    monkeypatch.setattr(orchestrator, "EvidenceRecollector", FakeEvidenceRecollector, raising=False)

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "input": {"case_filter": ""},
        }
    )

    trace_path = run_dir / "cases" / "TC_recollect" / "evidence_recollection_trace.json"
    trace = json.loads(trace_path.read_text(encoding="utf-8"))
    assert [item["attempt"] for item in trace] == [1, 2]
    assert trace[0]["goal_id"] == "g_order"
    assert trace[1]["action"] == "visibility_adjustment"
    assert all(item["action_risk_level"] == "low" for item in trace)
    execution_trace = json.loads(
        (run_dir / "cases" / "TC_recollect" / "execution_trace.json").read_text(encoding="utf-8")
    )
    assert [item["trace_type"] for item in execution_trace] == ["evidence_recollection", "evidence_recollection"]
    assert execution_trace[0]["action_id"] == "g_order"
    assert execution_trace[0]["planned_target"] == "verification_evidence_gap"
    assert execution_trace[0]["normalized_target"] == "refresh_current_screen_evidence"
    assert execution_trace[1]["normalized_target"] == "visibility_adjustment"
    assert execution_trace[0]["correction_step"] == {
        "type": "evidence_recollection",
        "attempt": 1,
        "trigger": "verification_evidence_gap",
    }


def test_run_offline_updates_state_graph_when_enabled(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_state_graph",
        internal_id="TC_state_graph",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="跳转校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    class FakeDeviceWorkflow:
        def execute_plan(self, plan, case_dir):
            before_summary = case_dir / "element_summaries" / "001_before_a1.json"
            after_summary = case_dir / "element_summaries" / "001_after_a1.json"
            before_summary.parent.mkdir(parents=True, exist_ok=True)
            before_summary.write_text(
                json.dumps(
                    {
                        "status": "success",
                        "visible_texts": ["行情", "更多"],
                        "elements": [{"text": "行情"}, {"text": "更多"}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            after_summary.write_text(
                json.dumps(
                    {
                        "status": "success",
                        "visible_texts": ["国内指数"],
                        "elements": [{"text": "国内指数"}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            return [
                ActionResult(
                    action_id="a1",
                    status=ActionStatus.SUCCESS,
                    locator_level="poco_text",
                    target_element={"target": "更多"},
                    before_screenshot="screenshots/001_before_a1.png",
                    after_screenshot="screenshots/001_after_a1.png",
                    element_summary_before="element_summaries/001_before_a1.json",
                    element_summary_after="element_summaries/001_after_a1.json",
                    notes=[],
                )
            ]

    monkeypatch.setattr(orchestrator, "DeviceWorkflow", lambda *args, **kwargs: FakeDeviceWorkflow())

    run_dir = orchestrator.run_offline(
        {
            "execution": {"mode": "device", "enable_state_graph": True},
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "input": {"case_filter": ""},
        }
    )

    graph = json.loads((run_dir / "state_graph.json").read_text(encoding="utf-8"))
    assert graph["pages"]
    assert graph["edges"] == [
        {
            "from": next(page for page in graph["pages"] if graph["pages"][page]["summary"] == "行情|更多"),
            "action": "a1",
            "to": next(page for page in graph["pages"] if graph["pages"][page]["summary"] == "国内指数"),
            "case_id": "TC_state_graph",
        }
    ]


def test_run_offline_blocks_case_when_plan_exceeds_max_steps(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_step_budget",
        internal_id="TC_step_budget",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="步骤预算",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：进入国内指数并点击更多",
        parameters="",
        expected_result="进入国内指数列表页",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "execution": {"max_steps_per_case": 1},
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "input": {"case_filter": ""},
        }
    )

    action_results = json.loads(
        (run_dir / "cases" / "TC_step_budget" / "action_results.json").read_text(encoding="utf-8")
    )
    assert action_results == [
        {
            "action_id": "step_budget",
            "status": "blocked",
            "locator_level": "",
            "target_element": None,
            "before_screenshot": "",
            "after_screenshot": "",
            "element_summary_before": "",
            "element_summary_after": "",
            "notes": ["Execution plan has 4 actions, exceeding max_steps_per_case=1."],
            "execution_rationale_id": "",
        }
    ]


def test_run_offline_applies_case_param_overrides_before_planning_and_evidence(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_case_param",
        internal_id="TC_case_param",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="参数覆盖",
        priority="high",
        step_name="",
        precondition="",
        operation_description="输入股票代码",
        parameters='{"stock_code": "000001"}',
        expected_result="数据展示正确",
        original_fields={},
    )
    planned_parameters = {}

    class CapturingPlanner:
        def plan(self, planned_case):
            planned_parameters["value"] = planned_case.parameters
            return RealRuleBasedPlanner().plan(planned_case)

    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])
    monkeypatch.setattr(orchestrator, "RuleBasedPlanner", CapturingPlanner)

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "input": {"case_filter": "", "case_params": {"stock_code": "600519", "market": "沪深京"}},
        }
    )

    assert json.loads(planned_parameters["value"]) == {"stock_code": "600519", "market": "沪深京"}
    case_json = json.loads((run_dir / "cases" / "TC_case_param" / "case.json").read_text(encoding="utf-8"))
    assert json.loads(case_json["parameters"]) == {"stock_code": "600519", "market": "沪深京"}


def test_run_offline_passes_llm_client_to_verifier_when_enabled(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_llm_verifier",
        internal_id="TC_llm_verifier",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="数据校验",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="数据展示正确",
        original_fields={},
    )

    class FakeLLMClient:
        def __init__(self, provider=None, config=None, evidence_config=None, telemetry_sink=None):
            self.config = config
            self.evidence_config = evidence_config
            self.telemetry_sink = telemetry_sink

        def json_call(self, prompt, schema, **kwargs):
            return {
                "status": "success",
                "data": {
                    "observation_summary": "已采集页面证据，行情数据仍需人工核对。",
                    "manual_review_reason": "data_correctness",
                },
                "attempts": 1,
                "errors": [],
            }

    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])
    monkeypatch.setattr(orchestrator, "LLMClient", FakeLLMClient)

    run_dir = orchestrator.run_offline(
        {
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
            "verification": {"llm_preliminary_judgment": True},
            "input": {"case_filter": ""},
        }
    )

    verification = json.loads(
        (run_dir / "cases" / "TC_llm_verifier" / "verification_result.json").read_text(encoding="utf-8")
    )
    assert verification[0]["preliminary_status"] == "manual_required"
    assert verification[0]["structured_details"]["llm_status"] == "success"
    assert verification[0]["structured_details"]["llm_observation_summary"] == "已采集页面证据，行情数据仍需人工核对。"


def test_run_offline_initializes_empty_llm_usage_jsonl_when_llm_is_disabled(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_llm_disabled",
        internal_id="TC_llm_disabled",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="缓存指标产物",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="数据展示正确",
        original_fields={},
    )
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "llm": {"enabled": False},
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
        }
    )

    usage_path = run_dir / "llm_usage.jsonl"
    assert usage_path.exists()
    assert usage_path.read_text(encoding="utf-8") == ""


def test_run_offline_persists_shared_planning_and_verification_llm_usage(tmp_path, monkeypatch):
    case = NaturalLanguageTestCase(
        case_id="TC_llm_usage",
        internal_id="TC_llm_usage",
        row_number=2,
        business_module="股指",
        feature_module="国内指数",
        feature_item="",
        test_purpose="缓存指标审计",
        priority="high",
        step_name="",
        precondition="",
        operation_description="行情-股指：查看国内指数",
        parameters="",
        expected_result="数据展示正确",
        original_fields={},
    )

    class FakeLLMClient:
        def __init__(self, provider=None, config=None, evidence_config=None, telemetry_sink=None):
            self.telemetry_sink = telemetry_sink

        def json_call(self, prompt, schema, context=None, **kwargs):
            context = context or {}
            self.telemetry_sink(
                {
                    "timestamp": "2026-07-29T12:00:00+08:00",
                    "stage": context.get("stage", ""),
                    "case_id": context.get("case_id", ""),
                    "goal_id": context.get("goal_id", ""),
                    "attempt": 1,
                    "status": "success",
                    "model": "fake-model",
                    "request_id": f"req-{context.get('stage', '')}",
                    "usage_available": True,
                    "prompt_tokens": 10,
                    "prompt_cache_hit_tokens": 8,
                    "prompt_cache_miss_tokens": 2,
                    "completion_tokens": 2,
                    "total_tokens": 12,
                    "cache_hit_ratio": 0.8,
                    "error_type": "",
                    "error_message": "",
                }
            )
            if "actions" in schema.get("required", []):
                return {
                    "status": "success",
                    "data": {
                        "case_id": context["case_id"],
                        "understanding": "测试规划",
                        "preconditions": [],
                        "actions": [],
                        "verification_goals": [
                            {
                                "goal_id": "v1",
                                "claim": "页面数据可见",
                                "category": "data_correctness",
                                "expected_entities": ["国内指数"],
                                "evidence_priority": ["poco_tree", "ocr_text", "screenshot"],
                                "human_review_required": True,
                                "review_reason": "data_correctness",
                            }
                        ],
                        "manual_review_notes": [],
                        "interpretation_rationales": [],
                    },
                    "attempts": 1,
                    "errors": [],
                }
            return {
                "status": "success",
                "data": {
                    "observation_summary": "已采集页面证据。",
                    "manual_review_reason": "data_correctness",
                },
                "attempts": 1,
                "errors": [],
            }

    monkeypatch.setattr(orchestrator, "LLMClient", FakeLLMClient)
    monkeypatch.setattr(orchestrator, "_load_cases_or_dependency_case", lambda config: [case])

    run_dir = orchestrator.run_offline(
        {
            "llm": {
                "enabled": True,
                "use_for_planning": True,
                "use_for_verification": True,
            },
            "verification": {"llm_preliminary_judgment": True},
            "report": {"output_dir": str(tmp_path), "generate_html": False},
            "logs": {"enable_capture": False},
        }
    )

    records = [
        json.loads(line)
        for line in (run_dir / "llm_usage.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert {record["stage"] for record in records} == {"planning", "verification"}
    assert all(record["case_id"] == case.internal_id for record in records)
    assert all("prompt" not in record and "content" not in record for record in records)
