from pathlib import Path

from openpyxl import Workbook, load_workbook

from autoairtest.excel_writer import write_results_copy


def test_write_results_copy_updates_result_and_note_without_overwriting_source(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求测试报告"
    sheet.append(["TC_用例名称", "测试结果", "备注"])
    sheet.append(["TC_demo", "", "原备注"])
    workbook.save(source)

    output = write_results_copy(
        source,
        "需求测试报告",
        [
            {
                "case_id": "TC_demo",
                "run_status": "manual_required",
                "summary": "需要人工复核",
                "evidence_dir": "cases/TC_demo",
                "manual_review_reason": "data_correctness",
            }
        ],
        tmp_path / "result.xlsx",
    )

    assert output == tmp_path / "result.xlsx"
    original = load_workbook(source)["需求测试报告"]
    assert original.cell(row=2, column=2).value is None
    assert original.cell(row=2, column=3).value == "原备注"

    copied = load_workbook(output)["需求测试报告"]
    assert copied.cell(row=2, column=2).value == "manual_required"
    assert "需要人工复核" in copied.cell(row=2, column=3).value
    assert "cases/TC_demo" in copied.cell(row=2, column=3).value
    assert "data_correctness" in copied.cell(row=2, column=3).value


def test_write_results_copy_matches_internal_id_with_row_suffix(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求测试报告"
    sheet.append(["TC_用例名称", "测试结果", "备注"])
    sheet.append(["TC_dup", "", ""])
    workbook.save(source)

    output = write_results_copy(
        source,
        "需求测试报告",
        [{"case_id": "TC_dup__row_2", "run_status": "blocked", "summary": "执行被阻塞", "evidence_dir": ""}],
        tmp_path / "result.xlsx",
    )

    copied = load_workbook(output)["需求测试报告"]
    assert copied.cell(row=2, column=2).value == "blocked"
